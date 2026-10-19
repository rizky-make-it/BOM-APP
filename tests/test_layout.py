import os, sys
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
import openpyxl
from core.layout import compact_item_sheet, compact_items_sheet


def _vis(ws, a, b):
    return sum(1 for r in range(a, b + 1) if not ws.row_dimensions[r].hidden)


def test_item_sheet_follows_component_count():
    ws = openpyxl.Workbook().active
    compact_item_sheet(ws, 4, 7, 3)
    assert (_vis(ws, 18, 42), _vis(ws, 46, 60), _vis(ws, 64, 78), _vis(ws, 11, 14)) == (4, 7, 3, 1)


def test_zero_components_keeps_one_row():
    ws = openpyxl.Workbook().active
    compact_item_sheet(ws, 0, 0, 0)
    assert _vis(ws, 18, 42) == 1


def test_items_sheet():
    ws = openpyxl.Workbook().active
    compact_items_sheet(ws, 14)
    assert _vis(ws, 6, 53) == 14


def test_missing_price_highlight_rules():
    from core.layout import highlight_missing_item_sheet, highlight_missing_rates
    ws = openpyxl.Workbook().active
    highlight_missing_item_sheet(ws); highlight_missing_rates(ws)
    assert {str(r.sqref) for r in ws.conditional_formatting} == {"S18:T42", "S46:T60", "S64:T78", "C5:C400"}
