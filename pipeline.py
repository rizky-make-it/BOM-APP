"""
Full Automated BOM Pipeline for Factory Furniture Drawings.
Reads PDF drawing -> Gemini 3.8 Flash reasoning -> Factory rules derivation -> Golden Master Excel -> LibreOffice Recalc.
"""

import sys
import os
import json
import re
import base64
import sqlite3
import urllib.request
import subprocess
import shutil
from pathlib import Path
from typing import Dict, Any, List

import pymupdf
import openpyxl
from PIL import Image

try:
    from .config import DEFAULT_RATES, GOLDEN_MASTER_PATH
    from .core.rules_v2 import (
        derive_drawer_box,
        derive_door,
        derive_joint_deduction,
        derive_bed_anatomy,
        derive_buffet_casegoods,
        PART_CATALOG
    )
    from .models.item import ItemHeader, SolidPart, PanelPart, HardwarePart, BOMSpecification
    from .generators.single_bom import build_single_bom_workbook
except (ImportError, ValueError):
    from config import DEFAULT_RATES, GOLDEN_MASTER_PATH
    from core.rules_v2 import (
        derive_drawer_box,
        derive_door,
        derive_joint_deduction,
        derive_bed_anatomy,
        derive_buffet_casegoods,
        PART_CATALOG
    )
    from models.item import ItemHeader, SolidPart, PanelPart, HardwarePart, BOMSpecification
    from generators.single_bom import build_single_bom_workbook

SO_PATH = r"C:\Program Files\LibreOffice\program\soffice.exe"

def get_9router_key() -> str:
    db_path = os.path.join(os.environ.get('APPDATA', ''), '9router', 'db', 'data.sqlite')
    if os.path.exists(db_path):
        con = sqlite3.connect(f"file:{db_path.replace(chr(92), '/')}?mode=ro", uri=True)
        cur = con.cursor()
        cur.execute("SELECT key FROM apiKeys WHERE isActive=1 LIMIT 1")
        row = cur.fetchone()
        if row:
            return row[0]
    return os.environ.get('NINEROUTER_API_KEY', '')

def extract_pdf_data(pdf_path: str, work_dir: str) -> Dict[str, Any]:
    doc = pymupdf.open(pdf_path)
    text_facts = []
    image_paths = []
    
    dpi = 300
    scale = dpi / 72.0
    Image.MAX_IMAGE_PIXELS = None

    for pg_num in range(len(doc)):
        page = doc[pg_num]
        M = page.rotation_matrix
        
        # 1. Full page render
        full_png = os.path.join(work_dir, f"page_{pg_num}.png")
        page.get_pixmap(dpi=200).save(full_png)
        image_paths.append(full_png)
        
        # 2. Extract text tokens with visual coordinates
        tokens = []
        for b in page.get_text("dict")["blocks"]:
            for line in b.get("lines", []):
                txt = "".join(s["text"] for s in line["spans"]).strip()
                r = pymupdf.Rect(line["bbox"]) * M
                if txt:
                    tokens.append(f"{txt}@(x{int(r.x0)},y{int(r.y0)})")
        text_facts.append(f"PAGE {pg_num}: " + " ".join(tokens))
        
        # 3. High-res render for section crops
        hires_png = os.path.join(work_dir, f"hires_{pg_num}.png")
        page.get_pixmap(dpi=dpi).save(hires_png)
        im = Image.open(hires_png)
        
        if pg_num == 0:
            # Typical page 0: Front, Side, Top views
            crops = {
                "front": (im, (400, 170, 700, 520)),
                "side": (im, (200, 190, 420, 520)),
                "top": (im, (40, 30, 260, 220))
            }
        else:
            # Typical page 1: Sections C, B, A
            crops = {
                "secC": (im, (415, 40, 700, 215)),
                "secB": (im, (195, 215, 420, 490)),
                "secA": (im, (415, 215, 700, 490))
            }
            
        for name, (img_obj, (x0, y0, x1, y1)) in crops.items():
            crop_im = img_obj.crop((int(x0 * scale), int(y0 * scale), int(x1 * scale), int(y1 * scale)))
            crop_path = os.path.join(work_dir, f"crop_{name}.png")
            crop_im.save(crop_path)
            image_paths.append(crop_path)

    return {
        "text_facts": "\n".join(text_facts),
        "image_paths": image_paths
    }

def get_workbook_rates() -> Dict[str, float]:
    rates = {}
    wb = openpyxl.load_workbook(GOLDEN_MASTER_PATH, data_only=True)
    ws = wb["RATES"]
    for r in ws.iter_rows(values_only=True):
        if isinstance(r[0], str) and re.fullmatch(r'PLY-\d+-[A-Z]{3}', r[0]):
            vals = [x for x in r[1:4] if isinstance(x, (int, float))]
            if vals:
                rates[r[0]] = float(vals[0])
    return rates

def prompt_gemini_analysis(pdf_facts: str, image_paths: List[str], key: str) -> Dict[str, Any]:
    catalog_desc = "\n".join([f"- Part {k}: {v['name']} ({v['type']})" for k, v in PART_CATALOG.items()])
    
    prompt = f"""Anda adalah Chief Estimator & Furniture Engineering Expert untuk pabrik furnitur ekspor (standar Living Ikon).
Tugas Anda adalah membaca gambar kerja produksi furniture (lampiran gambar penuh & potongan resolusi tinggi) dan menyusun spesifikasi BOM yang 100% presisi.

ANGKA TEKS PDF (Otoritas absolut dimensi utama, titik koordinat tertera):
{pdf_facts}

ATURAN STRUKTURAL PABRIK (MANDAT RESMI V2.1):
1. KATALOG 25 PART RESMI:
{catalog_desc}

2. ATURAN CARCASS, RANGKA & SAMBUNGAN (STANDAR PABRIK RESMI V2.2):
- DEDUKSI SAMBUNGAN KAKI (JOINT DEDUCTION MANDATORY):
  * Rel horizontal yang masuk di antara 2 kaki sudut: Panjang Bersih = Bentang Luar - (Tebal Kaki Kiri + Tebal Kaki Kanan).
  * CONTOH: Jika lebar luar nakas 470 mm dan tebal kaki 40x40 mm, maka Panjang Rel = 470 - (40 + 40) = 390 mm! JANGAN gunakan ukuran luar 470 mm untuk rel!
- Komponen Carcass HANYA terdiri dari part fungsional riil:
  * Kaki (Leg): 4x sudut (45x45mm atau 40x40mm sesuai gambar)
  * Rel Panjang (front/back rail top and bottom): 4x (mengikat depan atas/bawah dan belakang atas/bawah)
  * Rel Samping Atas (top rail front to back): 2x
  * Stretcher Samping Bawah (Side stretcher): 2x
  * Partisi Tengah (center panel): 1x plywood 18mm diapit rangka penguat (frame @ center panel 2x solid 40x25mm)
  * Rangka Belakang (frame @ back panel): 1x tiang vertikal penguat tengah
  * Panel Top: 1x Plywood 18mm MSF
  * Panel Samping (Side panel): 2x Plywood 12mm MSF
  * Panel Belakang (Back panel): 2x Plywood 9mm RAW
  * JANGAN menambah rel/cleats bayangan yang tidak ada dalam daftar potongan resmi!

3. ATURAN STRUKTUR TEMPAT TIDUR (BED ANATOMY V2.2):
- PANJANG SIDE RAIL (Part 21):
  * Panjang Side Rail = Panjang Kasur + Celah Bersih (40 mm) + Tebal Tiang Headboard + Tebal Tiang Footboard = 2000 + 40 + 50 + 50 = 2140 mm!
  * JANGAN mengambil ukuran matras saja (2040 mm) karena side rail harus mengikat kaki headboard ke kaki footboard!
- STRETCHER HEADBOARD & FOOTBOARD (Part 1):
  * Panjang Stretcher = Lebar Luar - 2 * Lebar Tiang Kaki Sudut (Contoh Queen: 1690 - 2*80 = 1530 mm; King: 1990 - 2*80 = 1830 mm).
- LARANG PENGGABUNGAN LABEL (STRICT 1:1 CALLOUT MANDATE):
  * Top stretcher, middle stretcher, bottom stretcher headboard, dan bottom stretcher footboard WAJIB ditulis pada baris terpisah dengan dimensi dan profil masing-masing.
  * DILARANG MERANGKUM menjadi satu nama generic "Bottom back stretcher"!
  * Komponen pembagi vertikal (Part 25 & 26: 741 mm dan 190 mm) WAJIB diekstrak utuh dan tidak boleh diabaikan!

4. ATURAN MUKA LACI & KOTAK LACI (MANDAT TERBARU V2.2):
- KOTAK LACI DALAM (Side, Inner Front, Back, Bottom, Rel DTC): JANGAN Anda hitung manual! Sistem akan menghitung otomatis menggunakan rumus paten pabrik (derive_drawer_box):
  * Lebar Kotak: W_front - 64 mm (celah kanan/kiri 20mm + dinding 24mm)
  * Tinggi Kotak: H_front - 30 mm (celah atas 20mm, celah bawah 10mm, seragam keliling)
  * Dasar Laci: (W_front - 54 mm) x (D - 7 mm)
  * Rel Laci: DTC undermount sesuai kedalaman D

5. ATURAN CASEGOODS / BUFFET / CREDENZA (STANDAR PABRIK V2.3):
- DEDUKSI SAMBUNGAN: Panjang rel apron dan stretcher = Overall L/W - (2 x tebal kaki sudut).
- KAKI TINGGI (> 200 mm): Bagian dasar bawah bodi TIDAK menggunakan lantai triplek utuh, melainkan menggunakan rangka rel kayu solid pengikat bodi (apron rails 40x25 mm dan rel tengah 45x36 mm).
- PINTU (DOOR): Tinggi daun pintu bersih = bukaan kabinet - celah lis atas/bawah 25 mm (contoh bukaan 525 mm -> daun pintu 500 mm). Jumlah pintu DIKUNCI strictly sesuai nama item (misal "2 Doors" = 2 pcs, DILARANG dobel).
- PARTISI & AMBALAN: Partisi vertikal masuk menjangkau bibir rebate pintu (kedalaman 410 mm). Ambalan berpintu diberi lis kayu solid depan (520 x 30 x 20 mm).
- PANEL BELAKANG BERSEKAT: Jika ada tiang vertikal belakang, panel belakang dipotong per kompartemen (3 bay terpisah), bukan 1 lembar utuh.

6. FINISHING 3-ZONA:
- exterior_m2: Top + Front + Sides luar
- interior_m2: Belakang luar + dalam laci keliling + rel samping + alas laci + belakang laci
- sealer_only_m2: Belakang dalam + samping dalam + top dalam

4. ATURAN KODE MATERIAL:
- Solid: '<no> / A'
- Panel: '<no> / B-<tebal>-<0X|1X|2X>' atau '<no> / C-<tebal>-<0X|1X|2X>'
- panel_type harus valid: PLY-18-MSF, PLY-12-MSF, PLY-9-MSF, PLY-6-MSF, PLY-18-MDF, PLY-12-RAW, PLY-9-RAW, PLY-6-RAW.

FORMAT OUTPUT WAJIB:
Keluarkan HANYA JSON valid (tanpa markdown pembuka/penutup, atau bungkus dalam ```json ... ```):
{{
  "header": {{
    "item_code": "...",
    "item_name": "...",
    "collection": "...",
    "overall_l": 0,
    "overall_w": 0,
    "overall_h": 0
  }},
  "drawer": {{
    "count": 0,
    "front_w": 0,
    "front_h": 0,
    "depth_D": 0,
    "handles_per_drawer": 0,
    "drawer_front_type": "standard_mdf | framed_inlay",
    "evidence": "..."
  }},
  "solid_parts": [
    {{"label": "3 / A", "component": "Leg", "l": 0, "w": 0, "t": 0, "qty": 0, "exposed": "Y", "tier_override": "B", "note": "..."}}
  ],
  "panel_parts": [
    {{"label": "7 / B-18-1X", "component": "Top", "panel_type": "PLY-18-MSF", "l": 0, "w": 0, "t": 0, "qty": 0, "exposed_faces": 1, "note": "..."}}
  ],
  "hardware_parts": [
    {{"label": "16 / HANDLE", "description": "...", "qty": 0, "unit_price": 0, "uom": "pcs", "note": "..."}}
  ],
  "reading_notes": [
    {{"parameter": "...", "value": "...", "source": "...", "confidence": "HIGH|MED|LOW"}}
  ],
  "assumptions_and_doubts": ["..."]
}}
"""

    content = [{"type": "text", "text": prompt}]
    for img_path in image_paths:
        with open(img_path, "rb") as f:
            b64 = base64.b64encode(f.read()).decode("utf-8")
            content.append({
                "type": "image_url",
                "image_url": {"url": f"data:image/png;base64,{b64}"}
            })

    body = {
        "model": "ag/gemini-3.8-flash-high",
        "max_tokens": 20000,
        "stream": True,
        "messages": [{"role": "user", "content": content}]
    }

    req = urllib.request.Request(
        "http://127.0.0.1:20128/v1/chat/completions",
        json.dumps(body).encode("utf-8"),
        {"Content-Type": "application/json", "Authorization": f"Bearer {key}"}
    )

    chunks = []
    with urllib.request.urlopen(req, timeout=360) as resp:
        for line in resp:
            line = line.decode("utf-8", "replace").strip()
            if line.startswith("data:") and not line.endswith("[DONE]"):
                try:
                    delta = json.loads(line[5:])["choices"][0]["delta"].get("content") or ""
                    chunks.append(delta)
                except Exception:
                    pass

    raw_text = "".join(chunks)
    m = re.search(r'\{.*\}', raw_text, re.S)
    if not m:
        raise ValueError(f"Failed to parse JSON from Gemini response: {raw_text[:500]}")
    return json.loads(m.group(0))

def run_pipeline(pdf_path: str, output_xlsx: str = None) -> Dict[str, Any]:
    work_dir = os.path.join(os.path.expanduser("~"), "bom_out", "_pipe")
    shutil.rmtree(work_dir, ignore_errors=True)
    os.makedirs(work_dir, exist_ok=True)

    print(f"[*] Processing Drawing: {Path(pdf_path).name}")
    print("[1/5] Extracting PDF text and rendering views...")
    pdf_info = extract_pdf_data(pdf_path, work_dir)

    key = get_9router_key()
    if not key:
        raise RuntimeError("9router API key not found. Ensure 9router is running and logged in.")

    print("[2/5] Running Gemini 3.8 Flash analysis via 9router...")
    gemini_res = prompt_gemini_analysis(pdf_info["text_facts"], pdf_info["image_paths"], key)

    h = gemini_res["header"]
    dr = gemini_res.get("drawer", {})
    item_code = h.get("item_code", "ITEM-01")
    import datetime
    now_str = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
    pdf_stem = Path(pdf_path).stem
    default_name = f"{pdf_stem} {now_str}.xlsx"
    output_xlsx = output_xlsx or os.path.join(os.path.expanduser("~"), "bom_out", default_name)

    print(f"[3/5] Applying Factory Construction Rules (Item: {item_code})...")
    
    # 1. Panels handling & Drawer Box integration
    panel_parts: List[PanelPart] = []
    drop_components = {"drawer side", "drawer inner front", "drawer back", "drawer bottom", "drawer box side"}
    
    for p in gemini_res.get("panel_parts", []):
        comp = p["component"].strip().lower()
        if comp in drop_components:
            continue
        panel_parts.append(PanelPart(
            label=p["label"],
            component=p["component"],
            panel_type=p["panel_type"],
            l=float(p["l"]),
            w=float(p["w"]),
            t=float(p["t"]),
            qty=int(p["qty"]),
            exposed_faces=int(p.get("exposed_faces", 1)),
            note=p.get("note", "")
        ))

    hardware_parts: List[HardwarePart] = []
    for p in gemini_res.get("hardware_parts", []):
        desc = (p.get("description", "") + p.get("label", "")).lower()
        if "slide" in desc or "rel" in desc:
            continue
        hardware_parts.append(HardwarePart(
            label=p["label"],
            description=p["description"],
            qty=float(p["qty"]),
            unit_price=float(p.get("unit_price", 0.0)),
            uom=p.get("uom", "pcs"),
            item_master_code="HARDWARE",
            vendor="Lokal",
            note=p.get("note", "")
        ))

    # Apply strict drawer box formulas if drawers exist
    drawer_count = int(dr.get("count", 0))
    if drawer_count > 0:
        fw = float(dr["front_w"])
        fh = float(dr["front_h"])
        cw = float(h["overall_w"])
        d_len = float(dr.get("depth_D") or 350)
        
        box = derive_drawer_box(fw, fh, cw, side_length=d_len, qty=drawer_count)
        
        panel_parts.append(PanelPart(
            label="5 / B-12-0X",
            component="Drawer side",
            panel_type=box["side"].panel_type,
            l=box["side"].l,
            w=box["side"].w,
            t=box["side"].t,
            qty=int(box["side"].qty),
            exposed_faces=0,
            note="Rumus Pabrik Paten (derive_drawer_box)"
        ))
        panel_parts.append(PanelPart(
            label="11 / B-12-0X",
            component="Drawer box front/back",
            panel_type=box["box_front_back"].panel_type,
            l=box["box_front_back"].l,
            w=box["box_front_back"].w,
            t=box["box_front_back"].t,
            qty=int(box["box_front_back"].qty),
            exposed_faces=0,
            note="Rumus Pabrik Paten: (W-64) x (H-30)"
        ))
        panel_parts.append(PanelPart(
            label="12 / B-6-0X",
            component="Drawer bottom",
            panel_type=box["bottom"].panel_type,
            l=box["bottom"].l,
            w=box["bottom"].w,
            t=box["bottom"].t,
            qty=drawer_count,
            exposed_faces=0,
            note="Rumus Pabrik Paten: (W-54) x (D-7)"
        ))
        
        hardware_parts.append(HardwarePart(
            label=f"15 / DTC-{int(d_len)}",
            description=box["slide"].description,
            qty=float(drawer_count),
            unit_price=float(box["slide"].unit_price),
            uom="set",
            item_master_code="HARDWARE",
            vendor="Sukey / DTC",
            note="Rumus Pabrik: DTC slide"
        ))

    solid_parts = []
    for p in gemini_res.get("solid_parts", []):
        p_comp = p["component"].strip().lower()
        p_l = float(p["l"])
        p_w = float(p["w"])
        p_t = float(p["t"])
        p_qty = int(p["qty"])
        p_note = p.get("note", "")

        # Safety Rule: Bed side rail must span full exterior post length (2140 mm)
        if "side rail" in p_comp and ("bed" in h.get("item_name", "").lower() or "b01" in item_code.lower() or "b02" in item_code.lower()):
            if 2000.0 <= p_l <= 2060.0:
                p_l = 2140.0
                p_note = (p_note + " | Auto-adjusted to full post exterior span 2140mm").strip(" |")

        # Safety Rule: Mirror outer frame miter 45 deg length = outer frame size
        if ("mirror" in h.get("item_name", "").lower() or "b71" in item_code.lower()) and ("support" in p_comp or "frame" in p_comp):
            if "vertical" in p_comp and 880.0 <= p_l <= 920.0:
                p_l = float(h.get("overall_h") or 1000.0)
                p_note = (p_note + " | Miter 45 deg outer length").strip(" |")

        solid_parts.append(SolidPart(
            label=p["label"],
            component=p["component"],
            l=p_l,
            w=p_w,
            t=p_t,
            qty=p_qty,
            material="MINDI",
            exposed=p.get("exposed", "Y"),
            curved="N",
            tier_override=p.get("tier_override", "B"),
            note=p_note
        ))

    header = ItemHeader(
        item_code=item_code,
        item_name=h.get("item_name", "FURNITURE ITEM"),
        collection=h.get("collection", "LIVING IKON"),
        overall_l=float(h["overall_l"]),
        overall_w=float(h["overall_w"]),
        overall_h=float(h["overall_h"]),
        order_qty=1,
        drawing_link=Path(pdf_path).name,
        finish_recipe="NC NATURAL",
        flatpack_yn="N",
        notes=f"Generated via 9router Gemini + Factory Rules Pipeline"
    )

    spec = BOMSpecification(header, solid_parts, panel_parts, hardware_parts)
    
    # Save raw spec JSON
    spec_json_path = os.path.join(os.path.expanduser("~"), "bom_out", f"{item_code}_spec.json")
    with open(spec_json_path, "w", encoding="utf-8") as f:
        json.dump({
            "header": header.__dict__,
            "drawer": dr,
            "solid_parts": [p.__dict__ for p in solid_parts],
            "panel_parts": [p.__dict__ for p in panel_parts],
            "hardware_parts": [p.__dict__ for p in hardware_parts],
            "reading_notes": gemini_res.get("reading_notes", []),
            "assumptions_and_doubts": gemini_res.get("assumptions_and_doubts", [])
        }, f, indent=2, ensure_ascii=False)

    print("[4/5] Generating Golden Master Workbook...")
    wb_rates = get_workbook_rates()
    # Add proxy rates
    wb_rates["PLY-18-MDF"] = 590418.04
    wb_rates["PLY-12-RAW"] = 215000.0
    wb_rates["PLY-6-RAW"] = 115000.0
    
    raw_xlsx = os.path.join(work_dir, "raw_bom.xlsx")
    build_single_bom_workbook(spec, raw_xlsx, custom_rates=wb_rates)

    print("[5/5] Recalculating with LibreOffice (Ground Truth Formulas)...")
    rc_dir = os.path.join(work_dir, "rc")
    os.makedirs(rc_dir, exist_ok=True)
    subprocess.run([SO_PATH, "--headless", "--convert-to", "xlsx", "--outdir", rc_dir, raw_xlsx], check=True, timeout=180)
    
    final_recalc_xlsx = os.path.join(rc_dir, "raw_bom.xlsx")
    shutil.copy(final_recalc_xlsx, output_xlsx)
    shutil.rmtree(work_dir, ignore_errors=True)

    # Read verified figures from final workbook
    wb_res = openpyxl.load_workbook(output_xlsx, data_only=True)[item_code]
    direct_cost = wb_res["D96"].value
    misc_cost = wb_res["D97"].value
    oh_cost = wb_res["D98"].value
    unit_cost = wb_res["D99"].value
    open_checks = wb_res["D102"].value

    res = {
        "output_file": output_xlsx,
        "spec_file": spec_json_path,
        "header": header.__dict__,
        "drawer": dr,
        "direct_cost": direct_cost,
        "misc_cost": misc_cost,
        "overhead_cost": oh_cost,
        "unit_cost": unit_cost,
        "open_checks": open_checks,
        "solid_parts": solid_parts,
        "panel_parts": panel_parts,
        "hardware_parts": hardware_parts,
        "assumptions": gemini_res.get("assumptions_and_doubts", [])
    }
    return res

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python -m bom_app.pipeline <drawing.pdf> [-o <out.xlsx>]")
        sys.exit(1)
    target_pdf = sys.argv[1]
    out_target = sys.argv[3] if len(sys.argv) > 3 and sys.argv[2] == "-o" else None
    result = run_pipeline(target_pdf, out_target)
    print("\n" + "=" * 65)
    print("  AUTOMATED BOM GENERATION SUCCESSFUL")
    print("=" * 65)
    print(f"File:       {result['output_file']}")
    print(f"Direct:     Rp {result['direct_cost']:,.2f}")
    print(f"Unit (HPP): Rp {result['unit_cost']:,.2f}")
    print(f"Checks:     {result['open_checks']} (0 = PASS)")
    print("=" * 65)
