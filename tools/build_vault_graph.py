"""Build the Obsidian knowledge graph for BOM-APP from real data (specs, rate card, git, tests, drawings).

Writes ONLY into <vault>/10-Auto/ (wiped and rebuilt each run) and <vault>/.obsidian/graph.json.
Hand-written notes (00-09 folders) are never touched. Run: python tools/build_vault_graph.py
"""
import ast
import datetime
import glob
import json
import os
import re
import shutil
import subprocess
import sys

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
VAULT = r"C:\Users\Admin\Documents\Hermes Vault"
AUTO = os.path.join(VAULT, "10-Auto")
sys.path.insert(0, ROOT)
os.chdir(ROOT)

from models.item import ItemHeader, SolidPart, PanelPart, HardwarePart, BOMSpecification  # noqa: E402
from core.costing_engine import calculate_bom_costing  # noqa: E402
from core.panel_rules import to_e0, apply_meeting_rule, role_of  # noqa: E402
from core.curves import apply_to_spec  # noqa: E402
from generators.multi_bom import get_workbook_rates  # noqa: E402

R = dict(get_workbook_rates())
R["FX_USD"] = 18040.0


def _full_rate_card():
    """Every value on RATES (formulas recalculated by LibreOffice), so notes show real prices."""
    import tempfile
    import openpyxl
    from config import GOLDEN_MASTER_PATH
    t = tempfile.mkdtemp()
    shutil.copy(GOLDEN_MASTER_PATH, os.path.join(t, "g.xlsx"))
    so = "C:/Program Files/LibreOffice/program/soffice.exe"
    try:
        subprocess.run([so, "--headless", "--convert-to", "xlsx", "--outdir", os.path.join(t, "o"), os.path.join(t, "g.xlsx")],
                       check=True, timeout=240, capture_output=True)
        ws = openpyxl.load_workbook(os.path.join(t, "o", "g.xlsx"), data_only=True)["RATES"]
    except Exception as e:
        print("LibreOffice recalc gagal, pakai nilai tersimpan:", e)
        ws = openpyxl.load_workbook(GOLDEN_MASTER_PATH, data_only=True)["RATES"]
    out = {}
    for r in ws.iter_rows(min_row=1, values_only=True):
        if isinstance(r[0], str) and isinstance(r[2], (int, float)):
            out[r[0]] = float(r[2])
    shutil.rmtree(t, ignore_errors=True)
    return out


for _k, _v in _full_rate_card().items():
    R.setdefault(_k, _v)
IDR = lambda x: f"{x:,.0f}".replace(",", ".")
safe = lambda s: re.sub(r'[\\/:*?"<>|#^\[\]]+', "-", str(s)).strip(" .-") or "x"
slug = lambda s: re.sub(r"[^A-Za-z0-9]+", "-", str(s)).strip("-").lower()[:40] or "x"
NOTES = {}  # relative path (no .md) -> text


def put(path, text):
    assert path not in NOTES, path
    NOTES[path] = text.strip() + "\n"


def fm(tags, **kv):
    wide = "item" in tags or "koleksi" in tags or "peta" in tags or "harga" in tags
    out = ["---", "tags: [" + ", ".join(tags) + "]", "auto: true"]
    if wide:
        out.append("cssClasses: [table-wide, row-alt]")
    out += [f"{k}: {v}" for k, v in kv.items()]
    return "\n".join(out + ["---", ""])


def link(path, alias=None):
    return f"[[{path}|{alias}]]" if alias else f"[[{path}]]"


# ------------------------------------------------------------------ load specs
def load(f):
    d = json.load(open(f, encoding="utf-8"))
    sp = BOMSpecification(ItemHeader(**d["header"]), [SolidPart(**x) for x in d["solid_parts"]],
                          [PanelPart(**x) for x in d["panel_parts"]], [HardwarePart(**x) for x in d["hardware_parts"]])
    for p in sp.panel_parts:
        p.panel_type = to_e0(p.panel_type)
        apply_meeting_rule(p)
    apply_to_spec(sp, 0.0)
    return sp


files = (sorted(glob.glob("data/sienna/*_spec.json")) + sorted(glob.glob("data/napier/lv433_*_spec.json"))
         + sorted(glob.glob("data/lv453_*_spec.json")))
items = []  # dicts
seen = set()
for f in files:
    try:
        sp = load(f)
    except Exception as e:  # keep going, record the failure
        print("skip", f, e)
        continue
    code = sp.header.item_code
    if code in seen:
        continue
    seen.add(code)
    try:
        cost = calculate_bom_costing(sp, R)
    except Exception as e:
        print("cost fail", code, e)
        cost = {}
    coll = "SIENNA" if "/sienna/" in f.replace("\\", "/") else "NAPIER" if "/napier/" in f.replace("\\", "/") else "LV453"
    status = "terverifikasi dari gambar" if code == "LV433-B41" else "draf (asumsi banyak)" if coll == "NAPIER" else \
        "arsip 14 Okt" if coll == "LV453" else "spec dari Excel/gambar Rizky"
    items.append(dict(code=code, sp=sp, cost=cost, file=f.replace("\\", "/"), coll=coll, status=status))

# ------------------------------------------------------------------ helpers over rate card
def hollow_used(w, t):
    a, b = sorted((float(w), float(t)))
    wall = float(R.get("HOLLOW_WALL", 0.8))
    pre = f"HOLLOW-{a:g}x{b:g}x"
    for x in sorted(float(k[len(pre):]) for k in R if k.startswith(pre)):
        if x >= wall - 1e-9:
            return f"{pre}{x:g}", float(R[f"{pre}{x:g}"])
    return f"{pre}{wall:g}", None


# ------------------------------------------------------------------ git history
log = subprocess.run(["git", "log", "--name-only", "--format=@@%h|%ad|%s", "--date=format:%Y-%m-%d %H:%M"],
                     capture_output=True, text=True, encoding="utf-8", errors="replace").stdout
commits = []
for blk in log.split("@@")[1:]:
    lines = blk.strip().splitlines()
    h, d, s = lines[0].split("|", 2)
    commits.append(dict(h=h, d=d, s=s, files=[l.strip() for l in lines[1:] if l.strip()]))
file2item = {i["file"]: i["code"] for i in items}
item_commits = {i["code"]: [] for i in items}
for c in commits:
    c["items"] = sorted({file2item[f] for f in c["files"] if f in file2item})
    for code in c["items"]:
        item_commits[code].append(c)

RULE_KEYS = [(r"meeting|E0|grade", "01-Aturan-BOM/Material-Panel-E0"), (r"veneer", "01-Aturan-BOM/Veneer"),
             (r"hollow|metal", "01-Aturan-BOM/Metal-Hollow"), (r"curve|R30|profil", "01-Aturan-BOM/Lengkung-vs-Profil"),
             (r"drawer|laci", "01-Aturan-BOM/Rumus-Laci-Pabrik"), (r"thickness|RULES", "01-Aturan-BOM/Ketebalan-Rules-vs-Gambar"),
             (r"rate|price|RATES", "01-Aturan-BOM/Rate-dan-Engine")]

# ------------------------------------------------------------------ questions
qtext = open(os.path.join(VAULT, "04-Keputusan", "Pertanyaan-Terbuka.md"), encoding="utf-8").read() if os.path.exists(
    os.path.join(VAULT, "04-Keputusan", "Pertanyaan-Terbuka.md")) else ""
questions = []
for l in qtext.splitlines():
    m = re.match(r"^(\d+)\.\s+(.*)", l)
    if m:
        questions.append((int(m.group(1)), m.group(2)))
q_items = {n: set() for n, _ in questions}
for n, t in questions:
    for c in re.findall(r"LV\d{3}-[A-Z]\d{2}", t):
        if c in seen:
            q_items[n].add(c)

# ------------------------------------------------------------------ per item / component
rate_users = {}   # rate note path -> list of component links
mat_users = {}
hw_users = {}


def is_divider(p):
    return role_of(p.component) == "divider"


item_links_by_coll = {}
for it in items:
    sp, c, code = it["sp"], it["cost"], it["code"]
    h = sp.header
    ipath = f"10-Auto/Item/{safe(code)}"
    comp_links, rate_links, rule_links = [], set(), set()
    n = 0

    def comp_note(kind, p, extra_rows, rate_path, rate_label):
        global n
        n += 1
        cname = getattr(p, "component", None) or p.label
        cpath = f"10-Auto/Komponen/{safe(code)}_{n:02d}_{slug(cname)}"
        rows = [f"| Jenis | {kind} |", f"| Komponen | {cname} |", f"| Label gambar | {getattr(p, 'label', '')} |"] + extra_rows
        body = fm(["bom", "komponen", "auto"], item=code, jenis=kind) + f"# {code} - {cname}\n\nItem: {link(ipath, code)}\n\n| Field | Isi |\n|---|---|\n" + "\n".join(rows) + "\n"
        if rate_path:
            body += f"\nKode harga: {link(rate_path, rate_label)}\n"
        note = getattr(p, "note", "") or ""
        if note:
            body += f"\n**Catatan di spec:** {note}\n"
        put(cpath, body)
        comp_links.append(f"- {link(cpath, f'{n:02d} {cname}')} ({kind})")
        return cpath

    for p in sp.solid_parts:
        extra = [f"| Ukuran P x L x T | {p.l:g} x {p.w:g} x {p.t:g} mm |", f"| Qty | {p.qty:g} |", f"| Bahan | {p.material} |"]
        rp = rl = None
        if p.material == "METAL":
            hc, hp = hollow_used(p.w, p.t)
            rp, rl = f"10-Auto/Harga/{safe(hc)}", hc
            extra.append(f"| Harga hollow | {('Rp ' + IDR(hp) + '/m') if hp else 'TIDAK ADA di daftar (Rp 0)'} |")
            if hp:
                extra.append(f"| Biaya komponen | Rp {IDR(p.l / 1000 * p.qty * hp)} |")
            rule_links.add("01-Aturan-BOM/Metal-Hollow")
            rate_users.setdefault(rp, {"code": hc, "price": hp, "unit": "IDR/m", "kind": "hollow", "comps": []})
        else:
            rp, rl = f"10-Auto/Harga/MATERIAL-{safe(p.material)}", p.material
            rate_users.setdefault(rp, {"code": p.material, "price": None, "unit": "IDR/m3 (tier kayu)", "kind": "kayu", "comps": []})
        if getattr(p, "curved", "") == "Y":
            rule_links.add("01-Aturan-BOM/Lengkung-vs-Profil")
        cp = comp_note("kayu/metal", p, extra, rp, rl)
        rate_users[rp]["comps"].append(link(cp, f"{code} {p.component}"))
        rate_links.add(rp)

    for p in sp.panel_parts:
        pt = p.panel_type
        price = R.get(pt)
        extra = [f"| Ukuran P x L x T | {p.l:g} x {p.w:g} x {p.t:g} mm |", f"| Qty | {p.qty:g} |", f"| Kode panel | {pt} |",
                 f"| Muka veneer | {int(p.exposed_faces or 0)} |", f"| Peran (aturan meeting) | {role_of(p.component) or 'tidak diatur meeting'} |"]
        rp = f"10-Auto/Harga/{safe(pt)}"
        rate_users.setdefault(rp, {"code": pt, "price": price, "unit": "IDR/lembar", "kind": "papan", "comps": []})
        if is_divider(p):
            rule_links.add("01-Aturan-BOM/Material-Panel-E0")
        cp = comp_note("panel", p, extra, rp, pt)
        rate_users[rp]["comps"].append(link(cp, f"{code} {p.component}"))
        rate_links.add(rp)
        rule_links.add("01-Aturan-BOM/Material-Panel-E0")
        if "MSF" in pt or "MDF" in pt.split("E0-")[-1]:
            rule_links.add("01-Aturan-BOM/Veneer")
        if re.search(r"drawer", p.component, re.I):
            rule_links.add("01-Aturan-BOM/Rumus-Laci-Pabrik")

    unpriced = []
    for p in sp.hardware_parts:
        code_h = safe(p.label)
        rp = f"10-Auto/Hardware/{code_h}"
        hw_users.setdefault(rp, {"label": p.label, "desc": p.description, "price": p.unit_price, "uom": p.uom, "comps": []})
        extra = [f"| Deskripsi | {p.description} |", f"| Qty | {p.qty:g} {p.uom} |",
                 f"| Harga satuan | {('Rp ' + IDR(p.unit_price)) if p.unit_price else 'KOSONG (Rp 0)'} |", f"| Vendor | {p.vendor} |"]
        cp = comp_note("hardware", p, extra, rp, p.label)
        hw_users[rp]["comps"].append(link(cp, f"{code} {p.label}"))
        rate_links.add(rp)
        if not p.unit_price:
            unpriced.append(p.label)

    # rules from commit history keywords
    for cm in item_commits[code]:
        for pat, rp_ in RULE_KEYS:
            if re.search(pat, cm["s"], re.I):
                rule_links.add(rp_)
    rule_links.add("01-Aturan-BOM/Ketebalan-Rules-vs-Gambar")
    rule_links.add("01-Aturan-BOM/Rate-dan-Engine")

    qs = [n_ for n_, t in questions if code in q_items[n_] or
          (re.search(r"divider|sekat", t, re.I) and any(is_divider(p) for p in sp.panel_parts)) or
          (re.search(r"handle|KD fitting|insert nut|harga kosong", t, re.I) and unpriced) or
          (re.search(r"hollow|brass", t, re.I) and any(p.material == "METAL" for p in sp.solid_parts))]
    for q in qs:
        q_items[q].add(code)

    cost_rows = ""
    if c:
        cost_rows = "\n".join(f"| {k} | Rp {IDR(c.get(v, 0))} |" for k, v in [
            ("Kayu + metal", "timber_cost"), ("Fastener", "fasteners_cost"), ("Panel", "panels_cost"), ("Finishing material", "fin_mat_cost"),
            ("Amplas material", "sand_mat_cost"), ("Amplas tenaga", "sand_lab_cost"), ("Finishing tenaga", "fin_lab_cost"),
            ("Hardware", "hardware_cost"), ("Packing", "pack_cost"), ("Misc", "misc_cost"), ("Overhead", "overhead_cost"), ("**HPP unit**", "unit_cost")])
    drawings = [pth for pth in glob.glob("data/**/*.pdf", recursive=True) if code.replace("-", ".") in os.path.basename(pth) or code.replace("-", "") in os.path.basename(pth).replace(".", "")]
    body = fm(["bom", "item", "auto", it["coll"].lower()], kode=code, koleksi=it["coll"], status=f'"{it["status"]}"', hpp_mesin=int(c.get("unit_cost", 0)) if c else 0)
    body += f"# {code} - {h.item_name}\n\nKoleksi: {link('10-Auto/Koleksi/' + it['coll'], it['coll'])} | Status: **{it['status']}** | Ukuran P{h.overall_l:g} x L{h.overall_w:g} x T{h.overall_h:g} mm\n\nSpec: `{it['file']}`\n\n"
    body += "## HPP (mesin hitung, aturan terbaru)\n\n| Bagian | Biaya |\n|---|---|\n" + cost_rows + "\n\n(Excel menghitung panel dengan nesting sendiri; selisih 5-10%, lihat [[04-Keputusan/Pertanyaan-Terbuka]].)\n"
    body += f"\n## Komponen ({len(comp_links)})\n\n" + "\n".join(comp_links) + "\n"
    body += "\n## Kode harga yang dipakai\n\n" + "\n".join(f"- {link(x, x.split('/')[-1])}" for x in sorted(rate_links)) + "\n"
    if unpriced:
        body += "\n## Harga kosong\n\n" + "\n".join(f"- {x}" for x in sorted(set(unpriced))) + "\n"
    body += "\n## Aturan yang berlaku\n\n" + "\n".join(f"- {link(x)}" for x in sorted(rule_links)) + "\n"
    if qs:
        body += "\n## Pertanyaan terbuka terkait\n\n" + "\n".join(f"- {link(f'10-Auto/Pertanyaan/Q{q:02d}', f'Q{q:02d}')}" for q in sorted(set(qs))) + "\n"
    if drawings:
        body += "\n## Gambar\n\n" + "\n".join(f"- {link('10-Auto/Gambar/' + safe(os.path.basename(d)[:-4]), os.path.basename(d))}" for d in drawings) + "\n"
    if item_commits[code]:
        body += "\n## Riwayat commit yang menyentuh spec\n\n" + "\n".join(f"- {link('10-Auto/Commit/' + cm['h'], cm['h'])} {cm['d']}: {cm['s'][:90]}" for cm in item_commits[code][:15]) + "\n"
    put(ipath, body)
    item_links_by_coll.setdefault(it["coll"], []).append((code, h.item_name, c.get("unit_cost", 0) if c else 0))

# ------------------------------------------------------------------ rate / material / hardware notes
for rp, d in rate_users.items():
    body = fm(["bom", "harga", "auto"], kode=safe(d["code"]), jenis=d["kind"])
    body += f"# {d['code']}\n\nJenis: {d['kind']} | Harga: **{('Rp ' + IDR(d['price']) + ' ' + d['unit']) if d['price'] else 'tidak ada di rate card / tergantung tier'}**\n\n"
    if d["kind"] == "hollow":
        body += "Sumber: `data/prices/besi_hollow_hitam_2026-10-19.json`; aturan: [[01-Aturan-BOM/Metal-Hollow]]\n\n"
    elif d["kind"] == "papan":
        body += "Aturan: [[01-Aturan-BOM/Material-Panel-E0]], [[01-Aturan-BOM/Veneer]]\n\n"
    body += f"## Dipakai oleh {len(d['comps'])} komponen\n\n" + "\n".join(f"- {x}" for x in d["comps"]) + "\n"
    put(rp, body)
for rp, d in hw_users.items():
    body = fm(["bom", "hardware", "auto"], kode=safe(d["label"])) + f"# {d['label']}\n\n{d['desc']}\n\nHarga satuan: **{('Rp ' + IDR(d['price'])) if d['price'] else 'KOSONG (Rp 0)'}** per {d['uom']}\n\n## Dipakai oleh {len(d['comps'])} komponen\n\n" + "\n".join(f"- {x}" for x in d["comps"]) + "\n"
    put(rp, body)

# ------------------------------------------------------------------ whole rate card (used or not)
hollow_all = sorted(k for k in R if str(k).startswith("HOLLOW-"))
board_all = sorted(k for k in R if str(k).startswith("LND-"))
wood_all = sorted(k for k in R if re.fullmatch(r"(MINDI|TEAK|MAHOGANY)[A-Z_]*", str(k)))
NL = chr(10)
for k in hollow_all + board_all + wood_all:
    rp = f"10-Auto/Harga/{safe(k)}"
    if rp in NOTES:
        continue
    kind = "hollow" if k.startswith("HOLLOW-") else "papan" if k.startswith("LND-") else "kayu"
    unit = "IDR/m" if kind == "hollow" else "IDR/lembar" if kind == "papan" else "IDR/m3"
    body = fm(["bom", "harga", "auto"], kode=safe(k), jenis=kind, dipakai="tidak")
    body += f"# {k}" + NL + NL + f"Jenis: {kind} | Harga: **Rp {IDR(float(R[k]))} {unit}** | belum dipakai item mana pun" + NL + NL
    body += {"hollow": "Sumber: `data/prices/besi_hollow_hitam_2026-10-19.json`; aturan: [[01-Aturan-BOM/Metal-Hollow]]",
             "papan": "Aturan: [[01-Aturan-BOM/Material-Panel-E0]], [[01-Aturan-BOM/Veneer]]",
             "kayu": "Aturan: [[01-Aturan-BOM/Rate-dan-Engine]]"}[kind] + NL
    put(rp, body)

# ------------------------------------------------------------------ collections
for coll, lst in item_links_by_coll.items():
    tot = sum(x[2] for x in lst)
    body = fm(["bom", "koleksi", "auto"], koleksi=coll) + f"# Koleksi {coll}\n\n{len(lst)} item | total HPP mesin Rp {IDR(tot)}\n\n| Item | Nama | HPP mesin |\n|---|---|---|\n"
    body += "\n".join(f"| {link('10-Auto/Item/' + safe(c_), c_)} | {n_} | Rp {IDR(v)} |" for c_, n_, v in lst) + "\n"
    put(f"10-Auto/Koleksi/{coll}", body)

# ------------------------------------------------------------------ questions
for n_, t in questions:
    its = sorted(q_items[n_])
    body = fm(["bom", "pertanyaan", "auto"], nomor=n_) + f"# Q{n_:02d}\n\n{t}\n\nSumber: [[04-Keputusan/Pertanyaan-Terbuka]]\n\n## Item terdampak ({len(its)})\n\n" + "\n".join(f"- {link('10-Auto/Item/' + safe(x), x)}" for x in its) + "\n"
    put(f"10-Auto/Pertanyaan/Q{n_:02d}", body)

# ------------------------------------------------------------------ commits
for cm in commits:
    body = fm(["bom", "commit", "auto"], hash=cm["h"], tanggal=f'"{cm["d"]}"') + f"# {cm['h']}\n\n{cm['d']}\n\n**{cm['s']}**\n\n"
    if cm["items"]:
        body += "## Item yang spec-nya berubah\n\n" + "\n".join(f"- {link('10-Auto/Item/' + safe(x), x)}" for x in cm["items"]) + "\n\n"
    rules = sorted({rp_ for pat, rp_ in RULE_KEYS if re.search(pat, cm["s"], re.I)})
    if rules:
        body += "## Aturan terkait\n\n" + "\n".join(f"- {link(x)}" for x in rules) + "\n\n"
    tests = sorted({os.path.basename(f)[:-3] for f in cm["files"] if f.startswith("tests/") and f.endswith(".py")})
    if tests:
        body += "## Tes yang diubah" + chr(10) + chr(10) + chr(10).join((f"- {link('10-Auto/Tes/' + t_, t_)}" if os.path.exists(f"tests/{t_}.py") else f"- `{t_}` (sudah diganti atau dihapus)") for t_ in tests) + chr(10) + chr(10)
    body += f"## File ({len(cm['files'])})\n\n" + "\n".join(f"- `{f}`" for f in cm["files"][:40]) + ("\n- ..." if len(cm["files"]) > 40 else "") + "\n"
    put(f"10-Auto/Commit/{cm['h']}", body)

# ------------------------------------------------------------------ tests
ntests = 0
for tf in sorted(glob.glob("tests/test_*.py")):
    name = os.path.basename(tf)[:-3]
    tree = ast.parse(open(tf, encoding="utf-8").read())
    fns = [(n_.name, ast.get_docstring(n_) or "") for n_ in tree.body if isinstance(n_, ast.FunctionDef) and n_.name.startswith("test_")]
    ntests += len(fns)
    related = [r_ for pat, r_ in RULE_KEYS if re.search(pat.split("|")[0], name, re.I)]
    body = fm(["bom", "tes", "auto"], file=f"tests/{name}.py", jumlah=len(fns)) + f"# {name}\n\n`tests/{name}.py` - {len(fns)} tes\n\n" + "\n".join(f"- `{n_}` {d_}" for n_, d_ in fns) + "\n"
    if related:
        body += "\n## Aturan yang dijaga\n\n" + "\n".join(f"- {link(x)}" for x in sorted(set(related))) + "\n"
    used = [c_["h"] for c_ in commits if f"tests/{name}.py" in c_["files"]]
    if used:
        body += "\n## Commit yang mengubah tes ini\n\n" + "\n".join(f"- {link('10-Auto/Commit/' + h_, h_)}" for h_ in used) + "\n"
    put(f"10-Auto/Tes/{name}", body)

# ------------------------------------------------------------------ drawings
ndraw = 0
for pth in sorted(glob.glob("data/**/*.pdf", recursive=True)):
    base = os.path.basename(pth)[:-4]
    m = re.search(r"(LV\d{3})[.\-]([A-Z]\d{2})", base)
    icode = f"{m.group(1)}-{m.group(2)}" if m else None
    body = fm(["bom", "gambar", "auto"], file=f'"{pth}"') + f"# {base}\n\n`{pth}` ({os.path.getsize(pth) // 1024} KB)\n\n"
    if icode and icode in seen:
        body += f"Item: {link('10-Auto/Item/' + safe(icode), icode)}\n"
    put(f"10-Auto/Gambar/{safe(base)}", body)
    ndraw += 1

# ------------------------------------------------------------------ hub
hub = fm(["bom", "peta", "auto"]) + "# Peta pengetahuan BOM-APP (dibuat otomatis)\n\n"
hub += f"Dibangun {datetime.datetime.now():%d %b %Y %H:%M} oleh `tools/build_vault_graph.py` dari spec JSON, rate card, git, tes, dan gambar. **Jangan diedit; akan ditimpa.** Catatan tulisan tangan ada di folder 00-09.\n\n"
hub += f"| Jenis | Jumlah |\n|---|---|\n| Item | {len(items)} |\n| Komponen | {sum(1 for k in NOTES if k.startswith('10-Auto/Komponen/'))} |\n| Kode harga (dipakai / seluruh rate card) | {len(rate_users)} / {len(rate_users) + len(hollow_all + board_all + wood_all)} |\n| Hardware | {len(hw_users)} |\n| Commit | {len(commits)} |\n| Tes (file / fungsi) | {len(glob.glob('tests/test_*.py'))} / {ntests} |\n| Gambar PDF | {ndraw} |\n| Pertanyaan terbuka | {len(questions)} |\n\n"
hub += "## Koleksi\n\n" + "\n".join(f"- {link('10-Auto/Koleksi/' + c_, c_)}" for c_ in sorted(item_links_by_coll)) + "\n\n## Aturan (tulisan tangan)\n\n- [[00-Index]]\n- [[01-Aturan-BOM/Material-Panel-E0]]\n- [[01-Aturan-BOM/Veneer]]\n- [[01-Aturan-BOM/Metal-Hollow]]\n- [[04-Keputusan/Keputusan-Owner]]\n- [[04-Keputusan/Pertanyaan-Terbuka]]\n"
put("10-Auto/00-Peta", hub)

# ------------------------------------------------------------------ write
if os.path.isdir(AUTO):
    shutil.rmtree(AUTO)
for pth, text in NOTES.items():
    full = os.path.join(VAULT, *pth.split("/")) + ".md"
    os.makedirs(os.path.dirname(full), exist_ok=True)
    open(full, "w", encoding="utf-8", newline="\n").write(text)

# graph colours by folder
gdir = os.path.join(VAULT, ".obsidian")
if os.path.isdir(gdir):
    cols = [("10-Auto/Item", 0xE5484D), ("10-Auto/Komponen", 0x8E8E93), ("10-Auto/Harga", 0xF5A524), ("10-Auto/Hardware", 0xF97316),
            ("10-Auto/Commit", 0x3B82F6), ("10-Auto/Tes", 0x22C55E), ("10-Auto/Gambar", 0xA855F7), ("10-Auto/Pertanyaan", 0xEC4899),
            ("10-Auto/Koleksi", 0xFACC15), ("01-Aturan-BOM", 0x14B8A6), ("04-Keputusan", 0xF43F5E)]
    g = {"collapse-filter": True, "search": "", "showTags": False, "showAttachments": False, "hideUnresolved": False, "showOrphans": True,
         "collapse-color-groups": False,
         "colorGroups": [{"query": f"path:{p}", "color": {"a": 1, "rgb": c_}} for p, c_ in cols],
         "collapse-display": True, "showArrow": False, "textFadeMultiplier": -1, "nodeSizeMultiplier": 1, "lineSizeMultiplier": 1,
         "collapse-forces": True, "centerStrength": 0.4, "repelStrength": 12, "linkStrength": 1, "linkDistance": 120, "scale": 0.5, "close": False}
    json.dump(g, open(os.path.join(gdir, "graph.json"), "w", encoding="utf-8"), indent=1)

# ------------------------------------------------------------------ verify links
allpaths = set()
for root, _, fs in os.walk(VAULT):
    if ".obsidian" in root:
        continue
    for f in fs:
        if f.endswith(".md"):
            allpaths.add(os.path.relpath(os.path.join(root, f), VAULT)[:-3].replace("\\", "/"))
names = {p.split("/")[-1] for p in allpaths}
bad = []
nlinks = 0
for p in sorted(allpaths):
    t = open(os.path.join(VAULT, *p.split("/")) + ".md", encoding="utf-8").read()
    for l in re.findall(r"\[\[([^\]|#]+)", t):
        nlinks += 1
        if l not in allpaths and l not in names:
            bad.append((p, l))
print(f"catatan total {len(allpaths)} (auto {len(NOTES)}) | tautan {nlinks} | rusak {len(bad)}")
print(bad[:8])
print({k: sum(1 for x in NOTES if x.startswith('10-Auto/' + k + '/')) for k in ['Item', 'Komponen', 'Harga', 'Hardware', 'Commit', 'Tes', 'Gambar', 'Pertanyaan', 'Koleksi']})
