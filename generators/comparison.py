"""
Multi-Item Side-by-Side BOM Comparison Generator.
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_3way_comparison(files: list, output_path: str, title: str = "3-WAY BOM COMPARISON"):
    # Generates side by side comparison from 2 or 3 existing BOM files
    if len(files) < 2:
        raise ValueError("Comparison requires at least 2 BOM files.")

    wb_comp = openpyxl.Workbook()
    ws = wb_comp.active
    ws.title = "3-WAY COMPARISON"
    ws.views.sheetView[0].showGridLines = True

    # Styling
    f_title = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
    f_sec = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    f_th = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
    f_bold = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
    f_reg = Font(name="Segoe UI", size=10, color="334155")
    f_italic = Font(name="Segoe UI", size=9, italic=True, color="64748B")
    f_price = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
    f_tot = Font(name="Segoe UI", size=11, bold=True, color="0F172A")

    fill_dark = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    fill_sec = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    fill_v1 = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    fill_v2 = PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid")
    fill_v3 = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_tot = PatternFill(start_color="FDE68A", end_color="FDE68A", fill_type="solid")

    border_soft = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    border_double = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='double', color='0F172A')
    )

    align_c = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_l = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
    align_r = Alignment(horizontal="right", vertical="center", indent=1)

    # Header Banner
    ws.row_dimensions[2].height = 36
    ws.merge_cells("B2:J2")
    ws["B2"] = title
    ws["B2"].font = f_title
    ws["B2"].fill = fill_dark
    ws["B2"].alignment = align_c

    # Read each file
    wbs = [openpyxl.load_workbook(f, data_only=True) for f in files]
    item_sheets = []
    for w in wbs:
        # Find item sheet
        target = None
        for s in w.sheetnames:
            if s not in ['README', 'ITEMS', 'TEMPLATE', 'RATES', 'LISTS', 'PART LIST REFERENCE']:
                target = w[s]
                break
        item_sheets.append(target or w.active)

    # Section I: Macro Summary
    R_SUM = 5
    ws.row_dimensions[R_SUM].height = 28
    ws.merge_cells(f"B{R_SUM}:J{R_SUM}")
    ws[f"B{R_SUM}"] = "I. EXECUTIVE COST ROLL-UP SUMMARY"
    ws[f"B{R_SUM}"].font = f_sec
    ws[f"B{R_SUM}"].fill = fill_sec
    ws[f"B{R_SUM}"].alignment = align_l

    # Write Headers
    ws.cell(6, 2, "COST ELEMENT").font = f_th
    ws.cell(6, 2).alignment = align_c
    ws.cell(6, 2).border = border_soft
    ws.cell(6, 2).fill = fill_dark
    ws.cell(6, 2).font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")

    fills = [fill_v1, fill_v2, fill_v3]
    for idx, (sh, fl) in enumerate(zip(item_sheets, fills)):
        col_idx = 3 + idx * 2
        name = sh.title
        ws.merge_cells(start_row=6, start_column=col_idx, end_row=6, end_column=col_idx+1)
        c = ws.cell(6, col_idx, f"{name}\nCOST (IDR)")
        c.font = f_th
        c.fill = fl
        c.alignment = align_c
        c.border = border_soft
        ws.cell(6, col_idx+1).border = border_soft

    cost_rows = [
        ("Solid Wood (Timber)", "D87"),
        ("Fasteners & Glue", "D88"),
        ("Plywood Panels", "D89"),
        ("Finishing & Sanding", "D90:D93"),
        ("Hardware & Bought-In", "D94"),
        ("Packaging Carton", "D95"),
        ("TOTAL PRODUCTION COST", "D96"),
        ("Factory Miscellaneous (4%)", "D97"),
        ("TOTAL UNIT COST (HPP IDR)", "D98")
    ]

    for r_idx, (cname, cell_ref) in enumerate(cost_rows, 7):
        is_tot = "TOTAL" in cname
        ws.row_dimensions[r_idx].height = 28
        cB = ws.cell(r_idx, 2, cname)
        cB.font = f_bold if is_tot else f_reg
        cB.border = border_double if "TOTAL UNIT COST" in cname else border_soft
        cB.alignment = align_l
        if is_tot: cB.fill = fill_tot

        for s_idx, sh in enumerate(item_sheets):
            col_idx = 3 + s_idx * 2
            if ":" in cell_ref:
                start_c, end_c = cell_ref.split(":")
                val = sum(float(sh[f"D{r}"].value or 0) for r in range(int(start_c[1:]), int(end_c[1:])+1))
            else:
                val = float(sh[cell_ref].value or 0)

            cVal = ws.cell(r_idx, col_idx, val)
            cVal.number_format = '"Rp "#,##0'
            cVal.font = f_price if not is_tot else f_tot
            cVal.alignment = align_r
            cVal.border = border_double if "TOTAL UNIT COST" in cname else border_soft
            if is_tot: cVal.fill = fill_tot

            unit_cost = float(sh["D98"].value or 1)
            cShare = ws.cell(r_idx, col_idx+1, val / unit_cost if unit_cost > 0 else 0)
            cShare.number_format = '0.0%'
            cShare.font = f_italic
            cShare.alignment = align_c
            cShare.border = border_double if "TOTAL UNIT COST" in cname else border_soft

    # Set generous column widths
    ws.column_dimensions['B'].width = 32
    for c_idx in range(3, 10):
        ws.column_dimensions[get_column_letter(c_idx)].width = 18

    # Copy sheets
    for sh in item_sheets:
        new_ws = wb_comp.create_sheet(title=f"BOM - {sh.title}")
        for r in range(1, min(sh.max_row+1, 115)):
            new_ws.row_dimensions[r].height = sh.row_dimensions[r].height or 22
            for c in range(1, min(sh.max_column+1, 26)):
                new_ws.cell(r, c).value = sh.cell(r, c).value

    wb_comp.save(output_path)
    return True
